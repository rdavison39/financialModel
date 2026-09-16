"""
Importer for Nesbitt Burns portfolio Excel reports.

Reads a Nesbitt Burns portfolio Excel snapshot and extracts the account,
snapshot timestamp, cash balances, and security holdings.
"""

import re
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook


@dataclass
class ImportedCash:
    """Cash balance imported from a brokerage snapshot."""

    currency: str
    amount: Decimal


@dataclass
class ImportedHolding:
    """Security holding imported from a brokerage snapshot."""

    symbol: str
    company_name: str
    quantity: Decimal
    average_cost: Decimal
    price: Decimal
    market_value: Decimal
    unrealized_gain: Decimal
    unrealized_gain_percent: Decimal
    daily_change: Decimal
    daily_change_percent: Decimal
    previous_close: Decimal
    currency: str


@dataclass
class ImportedAccount:
    """Complete Nesbitt Burns account snapshot imported from Excel."""

    account_number: str
    snapshot_date: datetime
    cash: list[ImportedCash]
    holdings: list[ImportedHolding]


class NesbittImporter:
    """Import a Nesbitt Burns portfolio report."""

    def __init__(self, file_path: str | Path) -> None:
        self.file_path = Path(file_path)

    @staticmethod
    def _decimal_or_zero(value: object) -> Decimal:
        """Convert a numeric spreadsheet value to Decimal.

        Brokerage exports can contain text placeholders in optional
        numeric columns. Treat blank/non-numeric optional values as zero.
        """
        if value is None:
            return Decimal("0")

        if isinstance(value, Decimal):
            return value

        if isinstance(value, (int, float)):
            return Decimal(str(value))

        text = str(value).strip()

        if not text:
            return Decimal("0")

        try:
            return Decimal(text.replace(",", "").replace("$", "").strip())
        except InvalidOperation:
            return Decimal("0")

    def import_file(self) -> ImportedAccount:
        """Import the workbook."""
        workbook = load_workbook(
            filename=self.file_path,
            data_only=True,
        )

        worksheet = workbook["Holdings"]

        account_number, snapshot_date = self._parse_header(
            worksheet["F1"].value
        )

        cash = self._parse_cash(worksheet)
        holdings = self._parse_holdings(worksheet)

        return ImportedAccount(
            account_number=account_number,
            snapshot_date=snapshot_date,
            cash=cash,
            holdings=holdings,
        )

    # -------------------------------------------------------------
    # Header
    # -------------------------------------------------------------

    def _parse_header(
        self,
        header: str | None,
    ) -> tuple[str, datetime]:
        """Parse account number and snapshot date."""

        if not header:
            raise ValueError(
                "Nesbitt Burns report header is missing."
            )

        header_text = str(header)

        account_match = re.search(
            r"account\s*#\s*(\d+)",
            header_text,
            re.IGNORECASE,
        )

        if account_match is None:
            raise ValueError(
                "Could not parse Nesbitt Burns account number: "
                f"{header_text}"
            )

        account_number = account_match.group(1)

        timestamp_match = re.search(
            r"as of\s+(.+)$",
            header_text,
            re.IGNORECASE,
        )

        if timestamp_match and timestamp_match.group(1).strip():
            timestamp_text = timestamp_match.group(1).strip()

            try:
                snapshot_date = datetime.fromisoformat(timestamp_text)
            except ValueError as exc:
                raise ValueError(
                    "Could not parse Nesbitt Burns "
                    f"report timestamp: {timestamp_text}"
                ) from exc

        else:
            snapshot_date = datetime.fromtimestamp(
                self.file_path.stat().st_mtime
            )

        return account_number, snapshot_date

    # -------------------------------------------------------------
    # Cash
    # -------------------------------------------------------------

    def _parse_cash(
        self,
        worksheet,
    ) -> list[ImportedCash]:
        """Parse CAD and USD cash balances."""

        cash: list[ImportedCash] = []

        for row in range(4, 7):
            currency = worksheet.cell(
                row=row,
                column=1,
            ).value

            amount = worksheet.cell(
                row=row,
                column=3,
            ).value

            if currency is None or amount is None:
                continue

            currency_text = str(currency).strip()

            if currency_text not in {"CAD", "USD"}:
                continue

            cash.append(
                ImportedCash(
                    currency=currency_text,
                    amount=Decimal(str(amount)),
                )
            )

        return cash

    # -------------------------------------------------------------
    # Holdings
    # -------------------------------------------------------------

    def _parse_holdings(
        self,
        worksheet,
    ) -> list[ImportedHolding]:
        """Parse security holdings."""

        holdings: list[ImportedHolding] = []

        for row in range(
            13,
            worksheet.max_row + 1,
        ):
            symbol = worksheet.cell(
                row=row,
                column=1,
            ).value

            description = worksheet.cell(
                row=row,
                column=2,
            ).value

            quantity = worksheet.cell(
                row=row,
                column=4,
            ).value

            average_cost = worksheet.cell(
                row=row,
                column=5,
            ).value

            price = worksheet.cell(
                row=row,
                column=7,
            ).value

            market_value = worksheet.cell(
                row=row,
                column=11,
            ).value

            currency = worksheet.cell(
                row=row,
                column=12,
            ).value

            unrealized_gain = worksheet.cell(
                row=row,
                column=14,
            ).value

            unrealized_gain_percent = worksheet.cell(
                row=row,
                column=16,
            ).value

            daily_change = worksheet.cell(
                row=row,
                column=24,
            ).value

            daily_change_percent = worksheet.cell(
                row=row,
                column=25,
            ).value

            previous_close = worksheet.cell(
                row=row,
                column=27,
            ).value

            if symbol is None:
                continue

            symbol_text = str(symbol).strip()

            if symbol_text in {
                "CANADIAN DOLLAR",
                "US DOLLAR",
            }:
                continue

            if (
                quantity is None
                or price is None
                or market_value is None
            ):
                continue

            company_name = (
                str(description).strip()
                if description is not None
                else symbol_text
            )

            currency_text = (
                str(currency).strip()
                if currency is not None
                else "CAD"
            )

            holdings.append(
                ImportedHolding(
                    symbol=symbol_text,
                    company_name=company_name,
                    quantity=Decimal(str(quantity)),
                    average_cost=self._decimal_or_zero(average_cost),
                    price=Decimal(str(price)),
                    market_value=Decimal(str(market_value)),
                    unrealized_gain=self._decimal_or_zero(
                        unrealized_gain
                    ),
                    unrealized_gain_percent=self._decimal_or_zero(
                        unrealized_gain_percent
                    ),
                    daily_change=self._decimal_or_zero(
                        daily_change
                    ),
                    daily_change_percent=self._decimal_or_zero(
                        daily_change_percent
                    ),
                    previous_close=self._decimal_or_zero(
                        previous_close
                    ),
                    currency=currency_text,
                )
            )

        return holdings
