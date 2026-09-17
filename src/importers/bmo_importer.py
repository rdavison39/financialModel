"""
BMO InvestorLine Excel importer.

Reads a BMO portfolio Excel snapshot and extracts the account,
snapshot timestamp, cash balances, and security holdings.
"""

from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from pathlib import Path
import re

from openpyxl import load_workbook


@dataclass
class ImportedCash:
    """Cash balance imported from a brokerage snapshot."""

    currency: str
    amount: Decimal


@dataclass
class ImportedHolding:
    """Security holding imported from a brokerage snapshot."""

    symbol: str
    company_name: str
    quantity: Decimal
    average_cost: Decimal
    price: Decimal
    market_value: Decimal
    unrealized_gain: Decimal
    unrealized_gain_percent: Decimal
    daily_change: Decimal
    daily_change_percent: Decimal
    previous_close: Decimal
    currency: str


@dataclass
class ImportedAccount:
    """Complete BMO account snapshot imported from Excel."""

    account_number: str
    snapshot_date: datetime
    cash: list[ImportedCash]
    holdings: list[ImportedHolding]


class BMOImporter:
    """Imports a BMO InvestorLine Excel portfolio report."""

    SHEET_NAME = "Holdings"

    def __init__(self, file_path: str | Path) -> None:
        """Initialize the importer."""
        self.file_path = Path(file_path)

    def import_file(self) -> ImportedAccount:
        """Read the BMO Excel file and return its imported data."""
        workbook = load_workbook(
            filename=self.file_path,
            data_only=True,
        )

        if self.SHEET_NAME not in workbook.sheetnames:
            raise ValueError(
                f"Expected worksheet '{self.SHEET_NAME}' "
                f"was not found."
            )

        worksheet = workbook[self.SHEET_NAME]

        account_number, snapshot_date = self._read_report_header(
            worksheet.cell(1, 6).value
        )

        cash = self._read_cash(worksheet)
        holdings = self._read_holdings(worksheet)

        return ImportedAccount(
            account_number=account_number,
            snapshot_date=snapshot_date,
            cash=cash,
            holdings=holdings,
        )

    def _read_report_header(
        self,
        header: object,
    ) -> tuple[str, datetime]:
        """Extract the account number and snapshot timestamp."""
        if not isinstance(header, str):
            raise ValueError("BMO report header is missing.")

        match = re.search(
            r"account\s*#\s*(\d+).*?as of\s+"
            r"(.+?)\s*\(Eastern Daylight Time\)",
            header,
            re.IGNORECASE,
        )

        if not match:
            raise ValueError(
                f"Could not parse BMO report header: {header}"
            )

        account_number = match.group(1)
        date_text = match.group(2).strip()

        snapshot_date = datetime.strptime(
            date_text,
            "%a %b %d %Y %H:%M:%S GMT-0400",
        )

        return account_number, snapshot_date

    def _read_cash(self, worksheet) -> list[ImportedCash]:
        """Read cash balances from the Cash Details section."""
        cash: list[ImportedCash] = []

        for row in range(4, 6):
            currency = worksheet.cell(row, 1).value
            amount = worksheet.cell(row, 3).value

            if currency is None or amount is None:
                continue

            currency_text = str(currency).strip()

            # BMO includes a summary row such as "Total (in CAD)".
            # Only actual CAD and USD cash balances are imported.
            if currency_text not in {"CAD", "USD"}:
                continue

            cash.append(
                ImportedCash(
                    currency=currency_text,
                    amount=Decimal(str(amount)),
                )
            )

        return cash

    @staticmethod
    def _parse_numeric_value(value: object) -> Decimal:
        """Parse a numeric value from an Excel cell.

        BMO sometimes returns values such as "65.39 C" or "13.27 U"
        for previous close, where the trailing letter identifies the
        currency. Only the numeric portion belongs in the database.
        """
        if value is None:
            return Decimal("0")

        if isinstance(value, (int, float, Decimal)):
            return Decimal(str(value))

        match = re.search(
            r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)",
            str(value),
        )

        if match is None:
            raise ValueError(
                f"Could not parse numeric value: {value!r}"
            )

        return Decimal(match.group(0))

    def _read_holdings(self, worksheet) -> list[ImportedHolding]:
        """Read security holdings from the Holding Details section."""
        holdings: list[ImportedHolding] = []

        for row in range(12, worksheet.max_row + 1):
            symbol = worksheet.cell(row, 1).value

            if not symbol:
                continue

            company_name = worksheet.cell(row, 2).value
            quantity = worksheet.cell(row, 3).value
            average_cost = worksheet.cell(row, 4).value
            price = worksheet.cell(row, 6).value
            market_value = worksheet.cell(row, 10).value
            currency = worksheet.cell(row, 11).value
            unrealized_gain = worksheet.cell(row, 12).value
            unrealized_gain_percent = worksheet.cell(row, 14).value
            daily_change = worksheet.cell(row, 22).value
            daily_change_percent = worksheet.cell(row, 23).value
            previous_close = worksheet.cell(row, 25).value

            if quantity is None or market_value is None:
                continue

            holdings.append(
                ImportedHolding(
                    symbol=str(symbol),
                    company_name=str(company_name or ""),
                    quantity=Decimal(str(quantity)),
                    average_cost=Decimal(str(average_cost or 0)),
                    price=Decimal(str(price or 0)),
                    market_value=Decimal(str(market_value)),
                    unrealized_gain=Decimal(str(unrealized_gain or 0)),
                    unrealized_gain_percent=Decimal(
                        str(unrealized_gain_percent or 0)
                    ),
                    daily_change=Decimal(str(daily_change or 0)),
                    daily_change_percent=Decimal(
                        str(daily_change_percent or 0)
                    ),
                    previous_close=self._parse_numeric_value(
                        previous_close
                    ),
                    currency=str(currency or "CAD"),
                )
            )

        return holdings